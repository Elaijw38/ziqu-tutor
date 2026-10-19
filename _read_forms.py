# -*- coding: utf-8 -*-
"""读附件2（申报表）与附件3（其他事项说明）的结构与要点。"""
import io
import os

from docx import Document

FILES = [
    r"C:\Users\32058\Downloads\附件2：郑州大学首届AI智能体创新应用大赛项目申报表.docx",
    r"C:\Users\32058\Downloads\附件3：其他事项说明.docx",
]
for p in FILES:
    print("\n" + "=" * 84)
    print(os.path.basename(p), " %.0f KB" % (os.path.getsize(p) / 1024))
    print("=" * 84)
    doc = Document(p)
    for i, para in enumerate(doc.paragraphs):
        t = para.text.strip()
        if t:
            print("  P%02d| %s" % (i, t[:220]))
    for ti, table in enumerate(doc.tables, 1):
        print("  ---- 表 %d：%d 行 × %d 列 ----" % (ti, len(table.rows), len(table.columns)))
        for ri, row in enumerate(table.rows):
            cells, seen = [], set()
            for c in row.cells:
                txt = c.text.strip().replace("\n", " ⏎ ")
                if txt and txt in seen:
                    continue
                seen.add(txt)
                cells.append(txt)
            line = " || ".join(cells)
            if line:
                print("    R%02d| %s" % (ri, line[:300]))
